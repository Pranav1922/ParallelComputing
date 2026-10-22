"""Build the Word versions of the CS15 report drafts from the Markdown sources in docs/report/.

    python3 docs/report/build_docx.py

Sources (unchanged): CS15_Integrated_Report_DRAFT.md, CS15_Individual_Technical_Note_DRAFT.md,
CITATION_VERIFICATION.md, references.bib. Output: docs/report/word/*.docx. Figures are PNG renderings (qlmanage, 2400
px) of the accepted PDF figures in results/figures/<campaign>/, copied to docs/report/word/figures/ together with
their source CSVs. Requires python-docx (1.2.0 used).
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
OUT = HERE / "word"
FIG = OUT / "figures"
S01, S03, S02 = ("20261017T044734Z-s01_final-430d8a10", "20261017T083329Z-s03_final-b8e5be9a",
                 "20261017T115054Z-s02_final-898ee9e1")
MD = "20261018T015324Z-omp_exec_final-9c4c1f9b"  # measured C++/OpenMP replay (milestone M-D)
TEXT_W = 17.0  # cm: A4 width 21 cm minus 2 cm margins

# figure id -> (png name, campaign, source stem, caption)
FIGS = {
    "1": ("fig01_S01_makespan_vs_size", S01, "g1_makespan_vs_size",
          "S01 — DAG size vs simulated makespan / lower bound, test_id, one panel per workload class × processor "
          "count, median and IQR over the 10 instances per cell (seed mean per instance); six methods shown. Report "
          "view drawn from the saved plotted-summary CSV (subset saved as "
          "docs/report/word/figures/derived/fig01_S01_makespan_vs_size_report_subset.csv); the full figure with "
          "all methods and unseen-workload cells is the accepted one."),
    "2a": ("fig02a_S01_opt_time_vs_size", S01, "g2_opt_time_vs_size",
           "S01 — Measured optimization / scheduling time per run (seconds, log scale, battery power) vs DAG size; "
           "median and IQR over runs."),
    "2b": ("fig02b_S01_convergence", S01, "g2_convergence",
           "S01 — Convergence: best-so-far simulated makespan / HEFT vs fitness evaluations, median and IQR over "
           "runs; lower row zoomed."),
    "3": ("fig03_S01_improvement_by_class", S01, "g3_improvement_by_class",
          "S01 — Paired per-instance improvement (%) in simulated makespan vs FCFS-EAT (top) and vs the per-instance "
          "best heuristic (bottom), by workload class and unseen-workload set."),
    "4": ("fig04_S03_feature_vs_policy", S03, "g7_feature_vs_policy",
          "S03 — Workload feature vs selected policy (ADAPTIVE, test splits): the two most important tree features, "
          "points coloured by selected heuristic, hollow = novelty-guard fallback; dashed = tree thresholds."),
    "5": ("fig05_S03_adaptive_vs_fixed", S03, "g8_adaptive_vs_fixed",
          "S03 — Adaptive vs fixed policies: simulated makespan / lower bound (left) and measured time per call "
          "(right, seconds, log; mean, median and decision part)."),
    "6": ("fig06_S03_regret_heatmap", S03, "g9_regret_heatmap",
          "S03 — Robustness across workloads: mean regret (%) vs the per-instance best heuristic, workload class / "
          "unseen-workload set × policy; red box = worst class per policy."),
    "7": ("fig07_S02_training_curves", S02, "g4_training_curves",
          "S02 — Training reward (batch mean −Cmax/LB, rolling 20 updates) and validation ρ = Cmax_RL / Cmax_HEFT vs "
          "updates for the three seeds and the ablation (validation chose the checkpoints)."),
    "8": ("fig08_S02_makespan_by_method", S02, "g5_makespan_by_method",
          "S02 — Makespan vs method: simulated makespan / lower bound on test_id and test_ood (box = quartiles, "
          "whiskers = 5th–95th percentile, triangle = mean)."),
    "9": ("fig09_S02_rho_by_group", S02, "g6_rho_by_group",
          "S02 — Generalization: ρ = Cmax_RL / Cmax_HEFT per instance by workload class and unseen-workload set "
          "(< 1 = RL better)."),
    "10": ("fig10_MD_measured_execution", MD, "{md1_speedup_vs_threads,md2_simulator_gap}",
           "Measured execution (C++/OpenMP, Apple M4, mains power; wall-clock seconds). Left: measured speedup "
           "T1 / Tp vs threads at 1 ms mean task, median and IQR over the 360 test_id instances (T1: all tasks on "
           "one thread); dotted = ideal. Right: measured / predicted time of the HEFT schedule by thread count and "
           "mean task size, raw (nominal µs per sim_tu) and calibration-adjusted (measured serial seconds per "
           "sim_tu); 1 = simulator exact. Report view drawn from the saved plotted CSVs (no values recomputed)."),
}
# which figures replace each "*Figures (...)" paragraph, and how they are laid out
FIG_BLOCKS = {
    "S01": [("single", "1", 9.5), ("pair", ("2a", 8.2), [("2b", 8.0)]), ("single", "3", 9.5)],
    "S03": [("single", "4", 6.0), ("pair", ("5", 8.2), [("6", 8.0)])],
    "S02": [("single", "7", 11.5), ("pair", ("8", 8.2), [("9", 8.0)])],
    "MD": [("single", "10", 15.0)],
}

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_field(run, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def shade(cell, hex_fill: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_fill)
    tc.append(sh)


def add_inline(par, text: str, size=None, italic=False, bold=False) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, mono = bold, italic, False
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, mono = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], not italic
        if "\u2011" in part:  # Word's own non-breaking hyphen element (U+2011 has no glyph in many fonts)
            pieces = part.split("\u2011")
            for k, piece in enumerate(pieces):
                rr = par.add_run(piece)
                rr.bold, rr.italic = b or None, i or None
                if size:
                    rr.font.size = Pt(size)
                if k < len(pieces) - 1:
                    rr._r.append(OxmlElement("w:noBreakHyphen"))
            continue
        r = par.add_run(part)
        r.bold, r.italic = b or None, i or None
        if mono:
            r.font.name = "Consolas"
            r.font.size = Pt((size or 10.5) - 1)
        elif size:
            r.font.size = Pt(size)


def setup(doc: Document, header_text: str, body_pt: float = 10.5) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(body_pt)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.08
    for name, size, before in (("Heading 1", 13, 10), ("Heading 2", 11.5, 8), ("Heading 3", 10.5, 6)):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
        h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(before), Pt(3)
        h.paragraph_format.keep_with_next = True
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    cap = doc.styles["Caption"]
    cap.font.size, cap.font.italic, cap.font.bold = Pt(8.5), True, False
    cap.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    cap.paragraph_format.space_after = Pt(6)
    hp = sec.header.paragraphs[0]
    hp.text = header_text
    hp.runs[0].font.size, hp.runs[0].font.color.rgb = Pt(8), RGBColor(0x66, 0x66, 0x66)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t, field in (("Page ", None), (None, "PAGE"), (" of ", None), (None, "NUMPAGES")):
        r = fp.add_run(t) if t else fp.add_run()
        r.font.size = Pt(8)
        if field:
            add_field(r, field)


def add_table(doc, rows, widths=None, font=8.5, caption=None, cant_split=False):
    if caption:
        doc.add_paragraph(caption, style="Caption").paragraph_format.keep_with_next = True
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        # each column at least as wide as its longest unbreakable token (so "6.2e-52" or "base" never wraps),
        # the remaining width shared in proportion to text length
        char_cm = 0.019 * font + 0.002
        token = [max(len(tok) for r in rows for tok in re.sub(r"[*`]", "", r[c]).split() or [""]) for c in range(ncol)]
        floor = [min(max(t, 4) * char_cm + 0.35, 6.0) for t in token]
        lens = [min(max(len(r[c]) for r in rows), 45) ** 0.8 for c in range(ncol)]
        spare = max(TEXT_W - sum(floor), 0)
        widths = [f + spare * l / sum(lens) for f, l in zip(floor, lens)]
        widths = [w * TEXT_W / sum(widths) for w in widths]
    for j, w in enumerate(widths):
        t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, txt.strip(), size=font, bold=(i == 0))
            if i == 0:
                shade(cell, "DCE6F0")
    if cant_split:  # a row never breaks across pages (used in the note, whose result rows are tall)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in t.rows[0].cells:  # and the header row stays with the first data row
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True
    tr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    tr.append(h)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def picture(container, fid: str, width_cm: float, number_prefix="Figure"):
    png, camp, stem, cap = FIGS[fid]
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / f"{png}.png"), width=Cm(width_cm))
    c = container.add_paragraph(style="Caption")
    add_inline(c, f"**{number_prefix} {fid}.** {cap} Source: results/figures/{camp}/{stem}.{{pdf,svg,csv}}.",
               size=8)


def figure_block(doc, key: str) -> None:
    for item in FIG_BLOCKS[key]:
        if item[0] == "single":
            picture(doc, item[1], item[2])
            continue
        _, (left, lw), rights = item
        t = doc.add_table(rows=1, cols=2)
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        trpr = t.rows[0]._tr.get_or_add_trPr()   # keep both pictures and their captions on one page
        trpr.append(OxmlElement("w:cantSplit"))
        lc, rc = t.cell(0, 0), t.cell(0, 1)
        lc.width, rc.width = Cm(lw + 0.3), Cm(TEXT_W - lw - 0.3)
        t.columns[0].width, t.columns[1].width = Cm(lw + 0.3), Cm(TEXT_W - lw - 0.3)
        lc.paragraphs[0].text = ""
        picture(lc, left, lw)
        for fid, w in rights:
            picture(rc, fid, w)
        for cell in (lc, rc):  # drop the empty first paragraph python-docx creates
            first = cell.paragraphs[0]
            if not first.text and not first.runs:
                first._element.getparent().remove(first._element)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)


def md_blocks(text: str):
    """Yield (kind, payload) blocks from the restricted Markdown used in the drafts."""
    lines, i = text.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip() or ln.strip() == "---":
            i += 1
            continue
        if ln.startswith("```"):
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            yield "code", "\n".join(lines[i + 1:j])
            i = j + 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            yield f"h{len(m.group(1))}", m.group(2)
            i += 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c for c in re.split(r"(?<!\\)\|", lines[i].strip())[1:-1]]
                if not all(re.fullmatch(r"\s*:?-{3,}:?\s*", c) for c in cells):
                    rows.append([c.replace("\\|", "|") for c in cells])
                i += 1
            yield "table", rows
            continue
        if ln.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip("> ").rstrip())
                i += 1
            yield "quote", " ".join(buf)
            continue
        m = re.match(r"^(\s*)(- |\d+\. )(.*)", ln)
        if m:
            items = []
            while i < len(lines):
                mm = re.match(r"^(\s*)(- |\d+\. )(.*)", lines[i])
                if mm:
                    items.append([mm.group(2).strip() != "-", mm.group(3)])
                elif lines[i].startswith("  ") and lines[i].strip() and items:   # wrapped continuation line
                    items[-1][1] += " " + lines[i].strip()
                else:
                    break
                i += 1
            yield "list", items
            continue
        buf = [ln.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||>|```|- |\d+\. |---)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        yield "para", " ".join(buf)


def add_list(doc, items, size=None):
    for numbered, txt in items:
        p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        p.paragraph_format.space_after = Pt(1)
        add_inline(p, txt, size=size)


def render(doc, text: str, figures=True, h_offset=0, table_font=8.5, skip_title_block=False, cant_split=False):
    appendix = False
    for kind, data in md_blocks(text):
        if kind.startswith("h"):
            level = int(kind[1]) + h_offset
            if level <= 0:
                continue
            appendix = "Appendix" in data
            doc.add_heading(re.sub(r"[*`]", "", data), level=min(level, 3))
        elif kind == "para":
            m = re.match(r"^\*Figures \((S0\d|MD)", data)
            if m and figures:
                figure_block(doc, m.group(1))
                continue
            if re.match(r"^Table \d+ —", data):
                p = doc.add_paragraph(style="Caption")
                add_inline(p, re.sub(r"^(Table \d+) —", r"**\1.**", data), size=8.5)
                p.paragraph_format.keep_with_next = True
                continue
            p = doc.add_paragraph()
            if re.match(r"^\[\d+\]", data):  # reference list entry: hanging indent
                p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.8), Cm(-0.8)
                p.paragraph_format.space_after = Pt(1)
                add_inline(p, data, size=8.5)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline(p, data)
        elif kind == "table":
            if appendix:
                add_table(doc, data, widths=[3.6, 13.4], font=7.5)
            else:
                add_table(doc, data, font=table_font, cant_split=cant_split)
        elif kind == "list":
            add_list(doc, data)
        elif kind == "quote":
            p = doc.add_paragraph()
            add_inline(p, data, size=9, italic=True)
        elif kind == "code":
            for ln in data.splitlines():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(ln)
                r.font.name, r.font.size = "Consolas", Pt(8)


def title_block(doc, title: str, rows, banner: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold, r.font.size = True, Pt(15)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(banner)
    r.italic, r.bold, r.font.size = True, True, Pt(9)
    r.font.color.rgb = RGBColor(0xA0, 0x10, 0x10)
    add_table(doc, rows, widths=[4.0, 13.0], font=9)


def build_report():
    src = (HERE / "CS15_Integrated_Report_DRAFT.md").read_text()
    title = re.search(r"## 1\. Title\s+\*\*(.+?)\*\*", src, re.S).group(1).replace("\n", " ")
    body = src[src.index("## 2. Abstract"):]
    doc = Document()
    setup(doc, "BCSE412L Parallel Computing — CS15 Integrated Batch Report — Batch [BATCH_ID — PENDING] — DRAFT")
    title_block(doc, title, [
        ["Field", "Value"],
        ["Course", "BCSE412L Parallel Computing, Fall Semester 2026-27, Slot C1+TC1"],
        ["Case study", "CS15 — AI-Based and Adaptive Scheduling: S01 Evolutionary task scheduling; S02 RL-based "
                       "scheduling; S03 Adaptive workload-aware scheduler"],
        ["Batch ID", "[BATCH_ID — PENDING CONFIRMATION]"],
        ["Student ID(s)", "[STUDENT_ID — PENDING CONFIRMATION] — single-member batch covering S01, S02 and S03, "
                          "subject to faculty approval"],
        ["Status", "Draft, 2026-10-18. Pending: faculty approval of the one-person scope and confirmation that "
                   "the added C++/OpenMP components meet the tooling requirement; identifiers; Batches 1–3 and 7 "
                   "findings (S03); full texts of references [6], [7], [10], [16] ([5] partial)"],
    ], "DRAFT — not submission-ready (see Status)")
    body = body.replace("## 2. Abstract", "## 2. Abstract", 1)
    body = re.sub(r"\*\(≈ ?\d+ words\)\*", "", body)
    render(doc, body, h_offset=-1)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT / "CS15_Integrated_Report_DRAFT.docx")


def build_note():
    src = (HERE / "CS15_Individual_Technical_Note_DRAFT.md").read_text()
    body = src[src.index("## 1. Individual contribution"):]
    doc = Document()
    setup(doc, "BCSE412L Parallel Computing — CS15 Individual Technical Note — [STUDENT_ID — PENDING] — DRAFT")
    title_block(doc, "CS15 Individual Member Technical Note: AI-Based and Adaptive Scheduling (S01, S02, S03)", [
        ["Field", "Value"],
        ["Student ID", "[STUDENT_ID — PENDING CONFIRMATION]"],
        ["Batch ID", "[BATCH_ID — PENDING CONFIRMATION]"],
        ["Sub-problems owned", "S01, S02, S03 (single-member batch; scope pending faculty approval). Whether one "
                               "combined note or one note per sub-problem is expected is to be confirmed."],
        ["Code / data", "schedlab/ (Python), cpp/ (C++ simulator, OpenMP executor); raw CSVs results/raw/; "
                        "provenance provenance/; batch report CS15_Integrated_Report_DRAFT.docx"],
    ], "DRAFT — not submission-ready")
    render(doc, body, h_offset=-1, table_font=8.5, cant_split=True)
    p = doc.add_paragraph()
    add_inline(p, "Supporting figure (same as Figure 9 of the batch report):")
    picture(doc, "9", 14.0)
    doc.save(OUT / "CS15_Individual_Technical_Note_DRAFT.docx")


def build_literature():
    rep = (HERE / "CS15_Integrated_Report_DRAFT.md").read_text()
    lit = rep[rep.index("### 3.2 Literature review"):rep.index("## 4. Problem statements")]
    refs = rep[rep.index("## 10. References"):]  # last section (Appendix A moved to TRACEABILITY.md)
    ver = (HERE / "CITATION_VERIFICATION.md").read_text()
    doc = Document()
    setup(doc, "BCSE412L Parallel Computing — CS15 Literature Review and References — DRAFT")
    title_block(doc, "CS15 Literature Review and Verified References", [
        ["Field", "Value"],
        ["Purpose", "Literature review of the CS15 batch report (§3.2–3.3) with the verified reference list and the "
                    "citation-verification table"],
        ["Batch ID", "[BATCH_ID — PENDING CONFIRMATION]"],
        ["Verification", "Bibliographic records checked for all 16 references (Semantic Scholar, Crossref, OpenAlex); "
                         "content: full text [1], [4], [8], [12], [13]; partial [2], [3], [5], [9]; abstract only "
                         "[6], [7], [10], [16]; metadata only [11], [14], [15] (see table)"],
    ], "DRAFT — companion to the batch report")
    render(doc, lit.replace("### 3.2 Literature review", "## Literature review")
                   .replace("### 3.3 Research gap", "## Research gap"), h_offset=0)
    render(doc, refs.replace("## 10. References", "## References"), h_offset=0)
    v = ver[ver.index("| # | Reference"):]
    doc.add_heading("Citation verification", level=1)
    render(doc, ver[ver.index("**Two kinds"):ver.index("| # |")], h_offset=0)
    # compact verification table: drop the long reference column (the list above has it)
    rows = [[c.strip() for c in ln.split("|")[1:-1]] for ln in v.splitlines() if ln.startswith("|")]
    rows = [r for r in rows if not set("".join(r)) <= set("-: ")]
    keep = [0, 2, 3, 4, 5, 7]   # #, DOI, bibliographic check, content check, recent, notes
    add_table(doc, [[r[k] for k in keep] for r in rows], widths=[0.9, 3.4, 2.4, 4.4, 1.5, 4.4], font=7)
    render(doc, v[v.index("**Manual requirement"):], h_offset=0)
    doc.save(OUT / "CS15_Literature_Review_DRAFT.docx")


def make_report_fig1():
    """Report-only view of S01 G1, drawn from the saved plotted-summary CSV (no recomputation): the 4
    in-distribution classes x m in {2, 4, 8}, six methods. The full G1 (with every unseen-workload cell) stays the
    accepted figure in results/figures/<campaign>/."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import pandas as pd
    src = ROOT / "results" / "figures" / S01 / "g1_makespan_vs_size.csv"
    d = pd.read_csv(src, float_precision="round_trip")
    classes, procs = ["regular", "irregular", "heterogeneous", "irregular_hetero"], [2, 4, 8]
    algos = ["FCFS-EAT", "MINMIN", "HEFT", "HILL-CLIMB", "GA", "GA-KEYS-EFT"]
    d = d[d.Workload.isin(classes) & d.Resource_Count.isin(procs) & d.Algorithm.isin(algos)]
    out = FIG / "derived"
    out.mkdir(parents=True, exist_ok=True)
    d.to_csv(out / "fig01_S01_makespan_vs_size_report_subset.csv", index=False)
    fig, axes = plt.subplots(len(classes), len(procs), figsize=(6.7, 6.4), sharex=True, sharey="row")
    colors = dict(zip(algos, ["#d62728", "#9467bd", "#1f77b4", "#2ca02c", "#ff7f0e", "#8c564b"]))
    for r, cls in enumerate(classes):
        for c, m in enumerate(procs):
            ax = axes[r][c]
            for i, a in enumerate(algos):
                g = d[(d.Workload == cls) & (d.Resource_Count == m) & (d.Algorithm == a)].sort_values("Problem_Size")
                x = g.Problem_Size * (1 + 0.025 * (i - 2.5))
                ax.errorbar(x, g["median"], yerr=[g["median"] - g.q25, g.q75 - g["median"]], marker="o", ms=2.5,
                            lw=0.9, capsize=1.5, color=colors[a], label=a)
            ax.axhline(1.0, color="grey", ls=":", lw=0.7)
            ax.set_xscale("log")
            ax.set_xticks([50, 100, 200], ["50", "100", "200"])
            ax.minorticks_off()
            ax.tick_params(labelsize=6.5)
            ax.set_title(f"{cls}, m = {m}", fontsize=7)
            if c == 0:
                ax.set_ylabel("Cmax / LB", fontsize=7)
            if r == len(classes) - 1:
                ax.set_xlabel("DAG size n (tasks)", fontsize=7)
    axes[0][-1].legend(fontsize=5.5, loc="upper right")
    fig.tight_layout(pad=0.4)
    for ext in ("png", "pdf", "svg"):
        fig.savefig(out / f"fig01_S01_makespan_vs_size_report.{ext}", dpi=300)
    plt.close(fig)
    shutil.copy2(out / "fig01_S01_makespan_vs_size_report.png", FIG / "fig01_S01_makespan_vs_size.png")


def make_report_fig10():
    """Report view of the measured replay, drawn from the saved plotted CSVs of the frozen M-D analysis
    (md1_speedup_vs_threads.csv, md2_simulator_gap.csv); nothing is recomputed from raw data."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import pandas as pd
    src = ROOT / "results" / "figures" / MD
    s = pd.read_csv(src / "md1_speedup_vs_threads.csv", float_precision="round_trip")
    g = pd.read_csv(src / "md2_simulator_gap.csv", float_precision="round_trip")
    s = s[s.us_per_tu == 10.0]
    fig, (a, b) = plt.subplots(1, 2, figsize=(6.7, 2.6), gridspec_kw=dict(width_ratios=[1, 1.5]))
    for meth in ["HEFT", "GA#s0", "RL-S0", "ADAPTIVE", "FCFS-EAT", "TASKS"]:
        d = s[s.method == meth].sort_values("m")
        a.errorbar(d.m, d["median"], yerr=[d["median"] - d.q25, d.q75 - d["median"]], marker="o", ms=2.5, lw=0.9,
                   capsize=1.5, label={"TASKS": "OpenMP tasks", "GA#s0": "GA (seed 0)"}.get(meth, meth))
    a.plot([2, 8], [2, 8], ":", color="grey", lw=0.7)
    a.set_xticks([2, 4, 8])
    a.set_xlabel("threads (m)", fontsize=7)
    a.set_ylabel("measured speedup T1 / Tp", fontsize=7)
    a.tick_params(labelsize=6.5)
    a.legend(fontsize=5.5, loc="upper left")
    pos, data, labels = [], [], []
    for i, m in enumerate([2, 4, 8]):
        for j, (gap, us) in enumerate([("G_raw", 10.0), ("G_raw", 0.1), ("G_adj", 10.0), ("G_adj", 0.1)]):
            pos.append(i * 5 + j)
            data.append(g[(g.gap == gap) & (g.m == m) & (g.us_per_tu == us)].value.to_numpy())
    bp = b.boxplot(data, positions=pos, widths=0.8, showfliers=False, patch_artist=True)
    colors = ["#9ecae1", "#3182bd", "#fdae6b", "#e6550d"] * 3
    for box, c in zip(bp["boxes"], colors):
        box.set_facecolor(c)
    b.set_xticks([1.5, 6.5, 11.5], ["m = 2", "m = 4", "m = 8"], fontsize=7)
    b.axhline(1, color="grey", ls=":", lw=0.7)
    b.set_ylabel("measured / predicted (HEFT)", fontsize=7)
    b.tick_params(labelsize=6.5)
    from matplotlib.patches import Patch
    b.legend([Patch(color=c) for c in colors[:4]], ["raw, 1 ms tasks", "raw, 10 µs tasks", "adjusted, 1 ms tasks",
             "adjusted, 10 µs tasks"], fontsize=5.5, loc="upper left")
    fig.tight_layout(pad=0.4)
    out = FIG / "derived"
    out.mkdir(parents=True, exist_ok=True)
    plt.rcParams["svg.hashsalt"] = "cs15-fig10"
    fig.savefig(out / "fig10_MD_measured_execution_report.png", dpi=300)
    fig.savefig(out / "fig10_MD_measured_execution_report.pdf", metadata={"CreationDate": None})
    fig.savefig(out / "fig10_MD_measured_execution_report.svg", metadata={"Date": None})
    plt.close(fig)
    shutil.copy2(out / "fig10_MD_measured_execution_report.png", FIG / "fig10_MD_measured_execution.png")


def copy_sources():
    """Editable figure sources and CSVs next to the Word files (manual: plots in editable/source form)."""
    for fid, (png, camp, stem, _) in FIGS.items():
        d = FIG / "source" / f"{png}"
        d.mkdir(parents=True, exist_ok=True)
        for one in stem.strip("{}").split(","):
            for f in (ROOT / "results" / "figures" / camp).glob(f"{one}*"):
                shutil.copy2(f, d / f.name)
    tdir = OUT / "tables"
    tdir.mkdir(parents=True, exist_ok=True)
    for camp in (S01, S03, S02, MD):
        for f in (ROOT / "results" / "derived" / camp).glob("*.csv"):
            if f.name not in ("master.csv", "per_instance.csv"):  # size; both regenerate from raw
                shutil.copy2(f, tdir / f"{camp.split('-')[1]}__{f.name}")


if __name__ == "__main__":
    copy_sources()
    make_report_fig1()
    make_report_fig10()
    build_report()
    build_note()
    build_literature()
    print("written:", *sorted(p.name for p in OUT.glob("*.docx")))
